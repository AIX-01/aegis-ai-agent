"""
보고서 템플릿 테스트 스크립트
HTML, DOCX, PPTX 템플릿에 샘플 데이터와 이미지를 채워서 생성합니다.
"""
import os
import io
import base64
from datetime import datetime
from pathlib import Path

# 프로젝트 루트 경로 (aegis-ai-agent)
PROJECT_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_DIR = PROJECT_ROOT / "templates" / "reports"
OUTPUT_DIR = PROJECT_ROOT / "mock_reports" / "test_output"

# 출력 디렉토리 생성
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def create_sample_frames(count=4):
    """테스트용 샘플 프레임 생성 (실제로는 CCTV 프레임)"""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        print("⚠️ Pillow 설치 필요: pip install Pillow")
        return []

    frames = []
    for i in range(count):
        # 640x360 회색 이미지 생성
        img = Image.new('RGB', (640, 360), color=(50 + i*20, 50 + i*20, 50 + i*20))
        draw = ImageDraw.Draw(img)

        # 텍스트 추가
        text = f"Frame {i+1}\n14:30:0{i}"
        draw.text((280, 160), text, fill=(255, 255, 255))

        # bytes로 변환
        buffer = io.BytesIO()
        img.save(buffer, format='JPEG', quality=80)
        frames.append(buffer.getvalue())

    return frames


def frames_to_base64_html(frames, max_width=150):
    """프레임들을 HTML img 태그로 변환 (4x2 그리드) - 테이블 기반으로 wkhtmltopdf 호환"""
    if not frames:
        return "<p>이미지 없음</p>"

    # 테이블 기반 레이아웃 (wkhtmltopdf 호환)
    html_parts = ['<table style="width: 100%; border-collapse: collapse; table-layout: fixed;">']

    for row in range(2):  # 2행
        html_parts.append('<tr>')
        for col in range(4):  # 4열
            idx = row * 4 + col
            if idx < len(frames):
                b64 = base64.b64encode(frames[idx]).decode('utf-8')
                html_parts.append(
                    f'<td style="width: 25%; padding: 4px; text-align: center; vertical-align: top;">'
                    f'<img src="data:image/jpeg;base64,{b64}" '
                    f'alt="Frame {idx+1}" '
                    f'style="width: 100%; max-width: 180px; border: 1px solid #ddd;">'
                    f'</td>'
                )
            else:
                html_parts.append('<td></td>')
        html_parts.append('</tr>')

    html_parts.append('</table>')
    return '\n'.join(html_parts)


# 샘플 프레임 생성 (8장)
sample_frames = create_sample_frames(8)

# HTML용 이미지 태그 생성
frames_html = frames_to_base64_html(sample_frames)

# 샘플 데이터 (실제 시스템과 동일한 형식)
sample_actions = [
    {"type": "field_action", "description": "보안팀 현장 출동 지시"},
    {"type": "emergency_call", "description": "112 긴급 신고 완료"},
    {"type": "field_action", "description": "CCTV PTZ 추적 활성화"},
]

# HTML용 actions (li 태그 포함)
actions_html = "\n".join([f"        <li>{a['description']}</li>" for a in sample_actions])

# DOCX/PPTX용 actions (일반 텍스트, 줄바꿈으로 구분)
actions_text = "\n".join([f"• {a['description']}" for a in sample_actions])

sample_data = {
    "occurred_at": "2026년 02월 09일 14:30:00",
    "event_type": "ASSAULT",
    "camera_location": "1층 로비",
    "camera_name": "CAM-001",
    "risk_level": "ABNORMAL",
    "risk_score": "0.85",
    "summary": "검은 후드티를 입은 중년 남성이 다른 남성을 주먹으로 폭행하고 있습니다. 피해자는 바닥에 쓰러져 있으며, 가해자는 계속해서 폭행을 가하고 있습니다.",
    "actions": actions_html,  # HTML용 (li 태그 포함)
    "actions_text": actions_text,  # DOCX/PPTX용 (일반 텍스트)
    "frames": frames_html,  # HTML용 이미지 (8장)
    "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
}

print(f"📸 샘플 프레임 생성: {len(sample_frames)}장")


def test_html_template():
    """HTML 템플릿 테스트"""
    template_path = TEMPLATE_DIR / "report_template.html"
    output_path = OUTPUT_DIR / "test_report.html"
    pdf_output_path = OUTPUT_DIR / "test_report.pdf"

    if not template_path.exists():
        print(f"❌ HTML 템플릿 없음: {template_path}")
        return False

    # 템플릿 읽기
    with open(template_path, "r", encoding="utf-8") as f:
        template = f.read()

    # 플레이스홀더 치환
    result = template
    for key, value in sample_data.items():
        result = result.replace(f"{{{{{key}}}}}", str(value))

    # HTML 저장
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(result)

    print(f"✅ HTML 생성 완료: {output_path}")

    # HTML → PDF 변환
    try:
        import pdfkit
        import platform

        # Windows에서 wkhtmltopdf 경로 지정
        config = None
        if platform.system() == "Windows":
            wkhtmltopdf_path = r"C:\Program Files\wkhtmltopdf\bin\wkhtmltopdf.exe"
            if os.path.exists(wkhtmltopdf_path):
                config = pdfkit.configuration(wkhtmltopdf=wkhtmltopdf_path)
                print(f"   wkhtmltopdf 경로: {wkhtmltopdf_path}")
            else:
                print(f"   wkhtmltopdf 파일 없음: {wkhtmltopdf_path}")

        options = {
            'encoding': 'UTF-8',
            'enable-local-file-access': ''
        }

        if config:
            pdfkit.from_file(str(output_path), str(pdf_output_path), configuration=config, options=options)
        else:
            pdfkit.from_file(str(output_path), str(pdf_output_path), options=options)

        print(f"✅ PDF 생성 완료: {pdf_output_path}")
    except OSError as e:
        print(f"⚠️ PDF 변환 OSError: {e}")
    except Exception as e:
        print(f"⚠️ PDF 변환 실패: {type(e).__name__}: {e}")

    return True


def test_docx_template():
    """DOCX 템플릿 테스트 (이미지 포함)"""
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError:
        print("❌ python-docx 설치 필요: pip install python-docx")
        return False

    template_path = TEMPLATE_DIR / "report_template.docx"
    output_path = OUTPUT_DIR / "test_report.docx"

    if not template_path.exists():
        print(f"❌ DOCX 템플릿 없음: {template_path}")
        return False

    # 템플릿 읽기
    doc = Document(template_path)

    # 단락에서 플레이스홀더 치환 (frames 제외)
    for para in doc.paragraphs:
        # paragraph 전체 텍스트 가져오기
        full_text = para.text

        # {{frames}} 처리
        if "{{frames}}" in full_text:
            # 기존 run 모두 제거
            for run in para.runs:
                run.text = ""
            # 이미지 4x2 그리드 삽입
            current_para = para
            for i, frame in enumerate(sample_frames[:8]):
                run = current_para.add_run()
                run.add_picture(io.BytesIO(frame), width=Inches(1.2))
                # 4개마다 새 paragraph로 줄바꿈
                if (i + 1) % 4 == 0 and i < 7:
                    new_para = doc.add_paragraph()
                    para._element.addnext(new_para._element)
                    current_para = new_para
            continue

        # 다른 플레이스홀더 처리
        for key, value in sample_data.items():
            if key == "frames":
                continue
            placeholder = f"{{{{{key}}}}}"
            if placeholder in full_text:
                # DOCX에서는 {{actions}}를 actions_text 값으로 치환
                if key == "actions":
                    value = sample_data["actions_text"]
                # 텍스트 치환
                new_text = full_text.replace(placeholder, str(value))
                for run in para.runs:
                    run.text = ""
                if para.runs:
                    para.runs[0].text = new_text
                else:
                    para.add_run(new_text)
                full_text = new_text

    # 테이블에서 플레이스홀더 치환
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    # paragraph 전체 텍스트 가져오기
                    full_text = para.text

                    # {{frames}} 처리 (단일 플레이스홀더)
                    if "{{frames}}" in full_text:
                        # 기존 run 모두 제거
                        for run in para.runs:
                            run.text = ""
                        # 이미지 4x2 그리드 삽입
                        current_para = para
                        for i, frame in enumerate(sample_frames[:8]):
                            run = current_para.add_run()
                            run.add_picture(io.BytesIO(frame), width=Inches(1.2))
                            # 4개마다 새 paragraph로 줄바꿈
                            if (i + 1) % 4 == 0 and i < 7:
                                current_para = cell.add_paragraph()
                        continue

                    # Frame 1 ~ Frame 8 개별 처리 (표 셀에 각각 있는 경우)
                    for i in range(1, 9):
                        frame_placeholder = f"Frame {i}"
                        if frame_placeholder in full_text and i <= len(sample_frames):
                            # 기존 run 모두 제거
                            for run in para.runs:
                                run.text = ""
                            # 해당 프레임 이미지 삽입 (고정 크기)
                            run = para.add_run()
                            run.add_picture(io.BytesIO(sample_frames[i-1]), width=Inches(1.8))
                            # 셀 가운데 정렬
                            from docx.enum.table import WD_TABLE_ALIGNMENT
                            from docx.enum.text import WD_ALIGN_PARAGRAPH
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            cell.vertical_alignment = WD_TABLE_ALIGNMENT.CENTER
                            break

                    # 다른 플레이스홀더 처리
                    for key, value in sample_data.items():
                        if key == "frames":
                            continue
                        placeholder = f"{{{{{key}}}}}"
                        if placeholder in full_text:
                            # DOCX에서는 {{actions}}를 actions_text 값으로 치환
                            if key == "actions":
                                value = sample_data["actions_text"]
                            # 기존 run 모두 제거하고 새 텍스트 설정
                            new_text = full_text.replace(placeholder, str(value))
                            for run in para.runs:
                                run.text = ""
                            if para.runs:
                                para.runs[0].text = new_text
                            else:
                                para.add_run(new_text)
                            full_text = new_text

    # 머리글/바닥글 플레이스홀더 치환
    for section in doc.sections:
        # 바닥글 처리
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer is not None:
                for para in footer.paragraphs:
                    full_text = ''.join([run.text for run in para.runs])
                    new_text = full_text
                    for key, value in sample_data.items():
                        if key == "frames":
                            continue
                        placeholder = f"{{{{{key}}}}}"
                        if placeholder in new_text:
                            if key == "actions":
                                value = sample_data["actions_text"]
                            new_text = new_text.replace(placeholder, str(value))
                    if new_text != full_text and para.runs:
                        para.runs[0].text = new_text
                        for run in para.runs[1:]:
                            run.text = ""

        # 머리글 처리
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header is not None:
                for para in header.paragraphs:
                    full_text = ''.join([run.text for run in para.runs])
                    new_text = full_text
                    for key, value in sample_data.items():
                        if key == "frames":
                            continue
                        placeholder = f"{{{{{key}}}}}"
                        if placeholder in new_text:
                            if key == "actions":
                                value = sample_data["actions_text"]
                            new_text = new_text.replace(placeholder, str(value))
                    if new_text != full_text and para.runs:
                        para.runs[0].text = new_text
                        for run in para.runs[1:]:
                            run.text = ""

    # 저장
    doc.save(output_path)
    print(f"✅ DOCX 생성 완료: {output_path}")
    return True


def test_pptx_template():
    """PPTX 템플릿 테스트 (이미지 포함)"""
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        print("❌ python-pptx 설치 필요: pip install python-pptx")
        return False

    template_path = TEMPLATE_DIR / "report_template.pptx"
    output_path = OUTPUT_DIR / "test_report.pptx"

    if not template_path.exists():
        print(f"❌ PPTX 템플릿 없음: {template_path}")
        return False

    # 템플릿 읽기
    prs = Presentation(template_path)

    # 모든 슬라이드의 텍스트 프레임에서 플레이스홀더 치환
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        for key, value in sample_data.items():
                            if key == "frames":
                                continue  # 이미지는 별도 처리
                            # PPTX에서는 {{actions}}를 actions_text 값으로 치환
                            if key == "actions":
                                if "{{actions}}" in run.text:
                                    run.text = run.text.replace("{{actions}}", sample_data["actions_text"])
                                continue
                            if f"{{{{{key}}}}}" in run.text:
                                run.text = run.text.replace(f"{{{{{key}}}}}", str(value))

            # 테이블 처리
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                for key, value in sample_data.items():
                                    if key == "frames":
                                        continue
                                    # PPTX에서는 {{actions}}를 actions_text 값으로 치환
                                    if key == "actions":
                                        if "{{actions}}" in run.text:
                                            run.text = run.text.replace("{{actions}}", sample_data["actions_text"])
                                        continue
                                    if f"{{{{{key}}}}}" in run.text:
                                        run.text = run.text.replace(f"{{{{{key}}}}}", str(value))

        # 슬라이드에 이미지 추가 ({{frames}} 플레이스홀더가 있는 슬라이드)
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                if "{{frames}}" in shape.text_frame.text:
                    # {{frames}} shape의 위치와 크기 가져오기
                    shape_left = shape.left
                    shape_top = shape.top
                    shape_width = shape.width
                    shape_height = shape.height

                    # 플레이스홀더 제거
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            run.text = run.text.replace("{{frames}}", "")

                    # 이미지 추가 (4x2 그리드 - 8장)
                    # shape 영역에 맞게 이미지 크기 계산
                    gap = Inches(0.08)
                    img_width = (shape_width - gap * 3) / 4  # 4열, 3개 간격
                    img_height = (shape_height - gap) / 2    # 2행, 1개 간격

                    # shape 왼쪽 상단을 기준으로 배치
                    left_start = shape_left
                    top_start = shape_top

                    for i, frame in enumerate(sample_frames[:8]):
                        row = i // 4
                        col = i % 4
                        left = left_start + col * (img_width + gap)
                        top = top_start + row * (img_height + gap)
                        slide.shapes.add_picture(
                            io.BytesIO(frame),
                            left, top,
                            width=img_width
                        )

    # 저장
    prs.save(output_path)
    print(f"✅ PPTX 생성 완료: {output_path}")
    return True


if __name__ == "__main__":
    print("=" * 60)
    print("보고서 템플릿 테스트")
    print("=" * 60)
    print(f"템플릿 디렉토리: {TEMPLATE_DIR}")
    print(f"출력 디렉토리: {OUTPUT_DIR}")
    print()

    # 템플릿 테스트
    test_html_template()
    test_docx_template()
    test_pptx_template()

    print()
    print("=" * 60)
    print(f"결과 파일 확인: {OUTPUT_DIR}")
    print("=" * 60)

